"""Excel export of the Invoices table (honouring the active filters)."""
from __future__ import annotations

import io
import json
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# Same pale orange as the web table, so a doubtful value stays visibly doubtful
# once the data leaves the app.
UNCERTAIN_FILL = PatternFill("solid", fgColor="FDE8D0")
EDITED_FILL = PatternFill("solid", fgColor="E3F0DC")

COLUMNS = [
    ("archive", "Archive", 26),
    ("filename", "Fichier", 42),
    ("supplier_name", "Nom du fournisseur", 28),
    ("invoice_number", "Numero de facture", 20),
    ("invoice_date", "Date de la facture", 16),
    ("total_amount", "Montant total", 14),
    ("currency", "Devise", 10),
    ("vat_amount", "Montant TVA", 14),
    ("extraction_notes", "Remarques", 46),
    ("error", "Erreur", 30),
]


def _value(row: Any, key: str) -> Any:
    try:
        return row[key]
    except (KeyError, IndexError, TypeError):
        return getattr(row, key, None)


def _json_list(row: Any, key: str) -> set:
    raw = _value(row, key)
    if not raw:
        return set()
    if isinstance(raw, (list, tuple, set)):
        return set(raw)
    try:
        return set(json.loads(raw))
    except (ValueError, TypeError):
        return set()


def build_workbook(rows: Sequence[Any]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Factures"

    sheet.append([label for _, label, _ in COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for row in rows:
        uncertain = _json_list(row, "uncertain_fields")
        edited = _json_list(row, "edited_fields")
        sheet.append([_value(row, key) for key, _, _ in COLUMNS])

        written = sheet[sheet.max_row]
        for index, (key, _, _) in enumerate(COLUMNS):
            if key in uncertain:
                written[index].fill = UNCERTAIN_FILL
            elif key in edited:
                written[index].fill = EDITED_FILL

    for index, (_, _, width) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:{}{}".format(get_column_letter(len(COLUMNS)), max(sheet.max_row, 1))

    legend = workbook.create_sheet("Legende")
    legend["A1"] = "Code couleur"
    legend["A1"].font = Font(bold=True)
    legend["A3"] = "Champ signale comme incertain par l'extraction (a verifier)"
    legend["A3"].fill = UNCERTAIN_FILL
    legend["A4"] = "Champ corrige a la main par un utilisateur"
    legend["A4"].fill = EDITED_FILL
    legend.column_dimensions["A"].width = 60

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
